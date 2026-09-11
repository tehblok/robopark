import csv
import io
from datetime import date, datetime

import pytest
from openpyxl import load_workbook
from sqlalchemy import event

from conftest import login_as, role_id_for
from robopark_api.models import (
    INVENTORY_INT64_MAX,
    InventoryCatalogComponent,
    InventoryCatalogPart,
    InventoryCount,
    InventoryCountLine,
    InventoryMovement,
    InventoryParkStock,
    InventoryReceipt,
    InventoryReceiptLine,
    Park,
    User,
    UserPark,
)
from robopark_api.security import hash_password
from robopark_api.services import inventory_exports


def _user(db, slug, username, parks=()):
    user = User(
        username=username,
        password_hash=hash_password("secret"),
        role_id=role_id_for(db, slug),
        access_status="approved",
        is_active=True,
    )
    db.add(user)
    db.flush()
    db.add_all(UserPark(user_id=user.id, park_id=park.id) for park in parks)
    db.commit()
    return user


def test_mechanic_can_export_assigned_park_as_utf8_bom_csv(
    client, db_session, seed_park_with_tracker
):
    mechanic = _user(db_session, "mechanic", "export-mechanic", [seed_park_with_tracker])
    component = InventoryCatalogComponent(
        name="=component",
        normalized_name="=component",
        created_by=mechanic.id,
        updated_by=mechanic.id,
    )
    db_session.add(component)
    db_session.flush()
    part = InventoryCatalogPart(
        component_id=component.id,
        name="+part",
        normalized_name="+part",
        article="=1+1",
        normalized_article="=1+1",
        created_by=mechanic.id,
        updated_by=mechanic.id,
    )
    db_session.add(part)
    db_session.flush()
    db_session.add(
        InventoryParkStock(
            park_id=seed_park_with_tracker.id,
            catalog_part_id=part.id,
            quantity=INVENTORY_INT64_MAX,
            minimum_quantity=2**40,
            location="@A-1",
            updated_by=mechanic.id,
        )
    )
    db_session.commit()
    login_as(client, mechanic.username, "secret")

    response = client.get(
        "/inventory/export",
        params={"park_id": seed_park_with_tracker.id, "format": "csv"},
    )

    assert response.status_code == 200, response.text
    assert response.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0] == [
        "ID парка",
        "Парк",
        "ID запчасти",
        "Компонента",
        "Запчасть",
        "Артикул",
        "Остаток",
        "Минимум",
        "Место",
        "Активна в парке",
    ]
    assert rows[1] == [
        str(seed_park_with_tracker.id),
        "Alpha",
        str(part.id),
        "'=component",
        "'+part",
        "'=1+1",
        str(INVENTORY_INT64_MAX),
        str(2**40),
        "'@A-1",
        "True",
    ]
    disposition = response.headers["content-disposition"]
    assert 'filename="inventory-park-' in disposition
    assert "filename*=UTF-8''%D0%A1%D0%BA%D0%BB%D0%B0%D0%B4-" in disposition


def test_export_role_scope_is_checked_before_inventory_queries(
    client, db_session, db_engine, seed_park_with_tracker
):
    foreign = Park(name="Foreign", tag="Foreign", is_active=True)
    db_session.add(foreign)
    db_session.commit()
    mechanic = _user(db_session, "mechanic", "export-scope-mechanic", [seed_park_with_tracker])
    operator = _user(db_session, "operator", "export-operator")
    admin = _user(db_session, "admin", "export-admin")

    statements = []

    def record_statement(_conn, _cursor, statement, _parameters, _context, _executemany):
        statements.append(statement.casefold())

    event.listen(db_engine, "before_cursor_execute", record_statement)
    try:
        login_as(client, mechanic.username, "secret")
        statements.clear()
        all_response = client.get("/inventory/export", params={"scope": "all", "format": "csv"})
        assert all_response.status_code == 403
        assert not any(
            table in statement
            for statement in statements
            for table in (
                "inventory_catalog_parts",
                "inventory_park_stocks",
                "inventory_movements",
                "inventory_receipts",
                "inventory_counts",
            )
        )

        statements.clear()
        foreign_response = client.get(
            "/inventory/export", params={"park_id": foreign.id, "format": "csv"}
        )
        assert foreign_response.status_code == 403
        assert not any("inventory_catalog_parts" in statement for statement in statements)

        login_as(client, operator.username, "secret")
        assert (
            client.get(
                "/inventory/export", params={"park_id": foreign.id, "format": "csv"}
            ).status_code
            == 200
        )
        assert (
            client.get("/inventory/export", params={"scope": "all", "format": "csv"}).status_code
            == 403
        )

        login_as(client, admin.username, "secret")
        assert (
            client.get("/inventory/export", params={"scope": "all", "format": "csv"}).status_code
            == 200
        )
    finally:
        event.remove(db_engine, "before_cursor_execute", record_statement)


def test_royal_xlsx_exports_four_operational_sheets_and_preserves_alias_history(
    client, db_session, seed_park_with_tracker
):
    foreign = Park(name="Beta", tag="Beta", is_active=True)
    db_session.add(foreign)
    db_session.commit()
    royal = _user(db_session, "royal", "export-royal")
    component = InventoryCatalogComponent(
        name="@Component",
        normalized_name="@component",
        created_by=royal.id,
        updated_by=royal.id,
    )
    db_session.add(component)
    db_session.flush()
    target = InventoryCatalogPart(
        component_id=component.id,
        name="-Target",
        normalized_name="-target",
        article="+TARGET",
        normalized_article="+target",
        created_by=royal.id,
        updated_by=royal.id,
    )
    db_session.add(target)
    db_session.flush()
    source = InventoryCatalogPart(
        component_id=component.id,
        name="Old source",
        normalized_name="old source",
        article="OLD",
        normalized_article="old",
        is_active=False,
        merged_into_part_id=target.id,
        created_by=royal.id,
        updated_by=royal.id,
    )
    db_session.add(source)
    db_session.flush()
    db_session.add_all(
        [
            InventoryParkStock(
                park_id=seed_park_with_tracker.id,
                catalog_part_id=target.id,
                quantity=INVENTORY_INT64_MAX,
                minimum_quantity=2**40,
                location="=A1",
                updated_by=royal.id,
            ),
            InventoryParkStock(
                park_id=foreign.id,
                catalog_part_id=target.id,
                quantity=7,
                minimum_quantity=3,
                location="B-4",
                updated_by=royal.id,
            ),
            InventoryMovement(
                park_id=seed_park_with_tracker.id,
                catalog_part_id=source.id,
                actor_user_id=royal.id,
                kind="receipt",
                delta=INVENTORY_INT64_MAX,
                balance_before=0,
                balance_after=INVENTORY_INT64_MAX,
                source_kind="receipt",
                source_id="receipt-line-1",
                note="@historic",
            ),
        ]
    )
    receipt = InventoryReceipt(
        park_id=seed_park_with_tracker.id,
        supplier="=Factory",
        document_number="+UPD-1",
        receipt_date=date(2026, 9, 11),
        comment="@accepted",
        status="posted",
        created_by=royal.id,
        posted_by=royal.id,
    )
    count = InventoryCount(
        park_id=foreign.id,
        name="-September",
        status="posted",
        created_by=royal.id,
        posted_by=royal.id,
    )
    db_session.add_all([receipt, count])
    db_session.flush()
    db_session.add_all(
        [
            InventoryReceiptLine(
                receipt_id=receipt.id,
                catalog_part_id=target.id,
                quantity=2**40,
                note="=box",
            ),
            InventoryCountLine(
                count_id=count.id,
                catalog_part_id=target.id,
                expected_quantity=4,
                actual_quantity=7,
                difference=3,
                comment="+checked",
            ),
        ]
    )
    db_session.commit()
    login_as(client, royal.username, "secret")

    response = client.get("/inventory/export", params={"scope": "all", "format": "xlsx"})

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    book = load_workbook(io.BytesIO(response.content), read_only=True, data_only=False)
    assert book.sheetnames == ["Остатки", "Движения", "Поставки", "Инвентаризации"]

    stock_rows = list(book["Остатки"].iter_rows(values_only=True))
    assert [row[0] for row in stock_rows[1:]] == [seed_park_with_tracker.id, foreign.id]
    assert {row[2] for row in stock_rows[1:]} == {target.id}
    assert stock_rows[1][3:6] == ("'@Component", "'-Target", "'+TARGET")
    assert stock_rows[1][6:9] == (str(INVENTORY_INT64_MAX), 2**40, "'=A1")

    movement_rows = list(book["Движения"].iter_rows(values_only=True))
    assert len(movement_rows) == 2
    assert movement_rows[1][3:5] == (source.id, target.id)
    assert movement_rows[1][9] == str(INVENTORY_INT64_MAX)
    assert movement_rows[1][13:16] == ("receipt-line-1", None, "'@historic")
    assert movement_rows[1][16] == royal.username

    receipt_rows = list(book["Поставки"].iter_rows(values_only=True))
    assert len(receipt_rows) == 2
    assert receipt_rows[1][4:8] == (
        "'=Factory",
        "'+UPD-1",
        datetime(2026, 9, 11),
        "'@accepted",
    )
    assert receipt_rows[1][8:10] == (royal.username, royal.username)
    assert receipt_rows[1][16:20] == ("'-Target", "'+TARGET", 2**40, "'=box")

    count_rows = list(book["Инвентаризации"].iter_rows(values_only=True))
    assert len(count_rows) == 2
    assert count_rows[1][3:7] == ("'-September", "posted", royal.username, royal.username)
    assert count_rows[1][13:19] == ("'-Target", "'+TARGET", 4, 7, 3, "'+checked")


@pytest.mark.parametrize("export_format", ["csv", "xlsx"])
def test_export_row_cap_returns_controlled_413(
    client, db_session, seed_park_with_tracker, monkeypatch, export_format
):
    royal = _user(db_session, "royal", f"cap-royal-{export_format}")
    component = InventoryCatalogComponent(
        name="Cap component",
        normalized_name="cap component",
        created_by=royal.id,
        updated_by=royal.id,
    )
    db_session.add(component)
    db_session.flush()
    db_session.add(
        InventoryCatalogPart(
            component_id=component.id,
            name="Cap part",
            normalized_name="cap part",
            article="CAP-1",
            normalized_article="cap-1",
            created_by=royal.id,
            updated_by=royal.id,
        )
    )
    db_session.commit()
    monkeypatch.setattr(inventory_exports, "MAX_EXPORT_ROWS", 0)
    login_as(client, royal.username, "secret")

    response = client.get(
        "/inventory/export",
        params={"park_id": seed_park_with_tracker.id, "format": export_format},
    )

    assert response.status_code == 413
    assert response.json() == {"detail": "inventory_export_too_large"}


def test_export_rejects_missing_ambiguous_and_invalid_parameters(
    client, db_session, seed_park_with_tracker
):
    royal = _user(db_session, "royal", "validation-royal")
    login_as(client, royal.username, "secret")

    missing = client.get("/inventory/export", params={"format": "csv"})
    ambiguous = client.get(
        "/inventory/export",
        params={"park_id": seed_park_with_tracker.id, "scope": "all", "format": "csv"},
    )
    invalid = client.get(
        "/inventory/export", params={"park_id": seed_park_with_tracker.id, "format": "pdf"}
    )
    invalid_scope = client.get(
        "/inventory/export",
        params={"park_id": seed_park_with_tracker.id, "scope": "selected", "format": "csv"},
    )

    assert missing.status_code == 400
    assert missing.json() == {"detail": "inventory_export_scope_required"}
    assert ambiguous.status_code == 400
    assert ambiguous.json() == {"detail": "inventory_export_scope_ambiguous"}
    assert invalid.status_code == 422
    assert invalid_scope.status_code == 422

from __future__ import annotations

import csv
import io
import tempfile
from collections.abc import Iterator, Mapping
from datetime import datetime

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select, true
from sqlalchemy.orm import Session, aliased

from robopark_api.models import (
    InventoryCatalogComponent,
    InventoryCatalogPart,
    InventoryCount,
    InventoryCountLine,
    InventoryMovement,
    InventoryParkStock,
    InventoryReceipt,
    InventoryReceiptLine,
    Park,
    User,
)
from robopark_api.services import inventory_access
from robopark_api.services.rbac import PERMISSION_INVENTORY_EXPORT, has_permission

MAX_EXPORT_ROWS = 100_000
MAX_XLSX_SIGNIFICANT_DIGITS = 15
XLSX_TEXT_CHUNK_SIZE = 32_766
XLSX_SPOOL_MEMORY_BYTES = 8 * 1024 * 1024
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

STOCK_HEADERS = (
    "ID парка",
    "Парк",
    "ID запчасти",
    "Компонента",
    "Запчасть",
    "Артикул",
    "Остаток",
    "Минимум",
    "Место",
    "Активна в парке",
)
MOVEMENT_HEADERS = (
    "ID движения",
    "ID парка",
    "Парк",
    "ID записи каталога",
    "Канонический ID запчасти",
    "Компонента",
    "Запчасть",
    "Артикул",
    "Тип движения",
    "Изменение",
    "Остаток до",
    "Остаток после",
    "Тип источника",
    "ID источника",
    "Задача",
    "Комментарий",
    "Исполнитель",
    "Дата",
)
RECEIPT_HEADERS = (
    "ID поставки",
    "ID парка",
    "Парк",
    "Статус",
    "Поставщик",
    "Номер документа",
    "Дата поставки",
    "Комментарий",
    "Создал",
    "Провёл",
    "Создано",
    "Проведено",
    "ID строки",
    "ID записи каталога",
    "Канонический ID запчасти",
    "Компонента",
    "Запчасть",
    "Артикул",
    "Количество",
    "Примечание",
)
COUNT_HEADERS = (
    "ID инвентаризации",
    "ID парка",
    "Парк",
    "Название",
    "Статус",
    "Создал",
    "Провёл",
    "Создано",
    "Проведено",
    "ID строки",
    "ID записи каталога",
    "Канонический ID запчасти",
    "Компонента",
    "Запчасть",
    "Артикул",
    "Учетное количество",
    "Фактическое количество",
    "Расхождение",
    "Комментарий",
)
CONTINUATION_HEADERS = ("Часть текста", "Всего частей")


def _park_ids(db: Session, user: User, *, park_id: int | None, all_parks: bool) -> list[int]:
    if not inventory_access.can_view_inventory(db, user) or not has_permission(
        db, user, PERMISSION_INVENTORY_EXPORT
    ):
        raise PermissionError("forbidden")
    if all_parks:
        if park_id is not None:
            raise ValueError("inventory_export_scope_ambiguous")
        if user.role not in {"admin", "royal"}:
            raise PermissionError("forbidden")
        return list(db.scalars(select(Park.id).order_by(Park.id)))
    if park_id is None:
        raise ValueError("inventory_export_scope_required")
    if user.role in {"admin", "royal"}:
        park = db.get(Park, park_id)
        if park is None:
            raise LookupError("park_not_found")
        return [park.id]
    inventory_access.require_park(db, user, park_id)
    return [park_id]


def _stock_statement(park_ids: list[int]):
    return (
        select(
            Park.id,
            Park.name,
            InventoryCatalogPart.id,
            InventoryCatalogComponent.name,
            InventoryCatalogPart.name,
            InventoryCatalogPart.article,
            func.coalesce(InventoryParkStock.quantity, 0),
            func.coalesce(InventoryParkStock.minimum_quantity, 0),
            InventoryParkStock.location,
            func.coalesce(InventoryParkStock.is_active, False),
        )
        .select_from(Park)
        .join(InventoryCatalogPart, true())
        .join(
            InventoryCatalogComponent,
            InventoryCatalogComponent.id == InventoryCatalogPart.component_id,
        )
        .outerjoin(
            InventoryParkStock,
            (InventoryParkStock.park_id == Park.id)
            & (InventoryParkStock.catalog_part_id == InventoryCatalogPart.id),
        )
        .where(
            Park.id.in_(park_ids),
            InventoryCatalogPart.is_active.is_(True),
            InventoryCatalogComponent.is_active.is_(True),
        )
        .order_by(Park.id, InventoryCatalogPart.normalized_name, InventoryCatalogPart.id)
    )


def _movement_statement(park_ids: list[int]):
    return (
        select(
            InventoryMovement.id,
            Park.id,
            Park.name,
            InventoryMovement.catalog_part_id,
            InventoryCatalogPart.merged_into_part_id,
            InventoryCatalogComponent.name,
            InventoryCatalogPart.name,
            InventoryCatalogPart.article,
            InventoryMovement.kind,
            InventoryMovement.delta,
            InventoryMovement.balance_before,
            InventoryMovement.balance_after,
            InventoryMovement.source_kind,
            InventoryMovement.source_id,
            InventoryMovement.issue_key,
            InventoryMovement.note,
            User.username,
            InventoryMovement.created_at,
        )
        .select_from(InventoryMovement)
        .join(Park, Park.id == InventoryMovement.park_id)
        .outerjoin(
            InventoryCatalogPart,
            InventoryCatalogPart.id == InventoryMovement.catalog_part_id,
        )
        .outerjoin(
            InventoryCatalogComponent,
            InventoryCatalogComponent.id == InventoryCatalogPart.component_id,
        )
        .join(User, User.id == InventoryMovement.actor_user_id)
        .where(InventoryMovement.park_id.in_(park_ids))
        .order_by(InventoryMovement.created_at, InventoryMovement.id)
    )


def _receipt_statement(park_ids: list[int]):
    creator = aliased(User)
    poster = aliased(User)
    return (
        select(
            InventoryReceipt.id,
            Park.id,
            Park.name,
            InventoryReceipt.status,
            InventoryReceipt.supplier,
            InventoryReceipt.document_number,
            InventoryReceipt.receipt_date,
            InventoryReceipt.comment,
            creator.username,
            poster.username,
            InventoryReceipt.created_at,
            InventoryReceipt.posted_at,
            InventoryReceiptLine.id,
            InventoryReceiptLine.catalog_part_id,
            InventoryCatalogPart.merged_into_part_id,
            InventoryCatalogComponent.name,
            InventoryCatalogPart.name,
            InventoryCatalogPart.article,
            InventoryReceiptLine.quantity,
            InventoryReceiptLine.note,
        )
        .select_from(InventoryReceipt)
        .join(Park, Park.id == InventoryReceipt.park_id)
        .join(creator, creator.id == InventoryReceipt.created_by)
        .outerjoin(poster, poster.id == InventoryReceipt.posted_by)
        .outerjoin(InventoryReceiptLine, InventoryReceiptLine.receipt_id == InventoryReceipt.id)
        .outerjoin(
            InventoryCatalogPart,
            InventoryCatalogPart.id == InventoryReceiptLine.catalog_part_id,
        )
        .outerjoin(
            InventoryCatalogComponent,
            InventoryCatalogComponent.id == InventoryCatalogPart.component_id,
        )
        .where(InventoryReceipt.park_id.in_(park_ids))
        .order_by(InventoryReceipt.created_at, InventoryReceipt.id, InventoryReceiptLine.id)
    )


def _count_statement(park_ids: list[int]):
    creator = aliased(User)
    poster = aliased(User)
    return (
        select(
            InventoryCount.id,
            Park.id,
            Park.name,
            InventoryCount.name,
            InventoryCount.status,
            creator.username,
            poster.username,
            InventoryCount.created_at,
            InventoryCount.posted_at,
            InventoryCountLine.id,
            InventoryCountLine.catalog_part_id,
            InventoryCatalogPart.merged_into_part_id,
            InventoryCatalogComponent.name,
            InventoryCatalogPart.name,
            InventoryCatalogPart.article,
            InventoryCountLine.expected_quantity,
            InventoryCountLine.actual_quantity,
            InventoryCountLine.difference,
            InventoryCountLine.comment,
        )
        .select_from(InventoryCount)
        .join(Park, Park.id == InventoryCount.park_id)
        .join(creator, creator.id == InventoryCount.created_by)
        .outerjoin(poster, poster.id == InventoryCount.posted_by)
        .outerjoin(InventoryCountLine, InventoryCountLine.count_id == InventoryCount.id)
        .outerjoin(
            InventoryCatalogPart,
            InventoryCatalogPart.id == InventoryCountLine.catalog_part_id,
        )
        .outerjoin(
            InventoryCatalogComponent,
            InventoryCatalogComponent.id == InventoryCatalogPart.component_id,
        )
        .where(InventoryCount.park_id.in_(park_ids))
        .order_by(InventoryCount.created_at, InventoryCount.id, InventoryCountLine.id)
    )


def _safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _csv_chunks(rows) -> Iterator[bytes]:
    try:
        yield b"\xef\xbb\xbf"
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer)
        writer.writerow(STOCK_HEADERS)
        yield buffer.getvalue().encode("utf-8")
        for row in rows:
            buffer.seek(0)
            buffer.truncate(0)
            writer.writerow(_safe_cell(value) for value in row)
            yield buffer.getvalue().encode("utf-8")
    finally:
        rows.close()


def _xlsx_cell(value):
    value = _safe_cell(value)
    if (
        isinstance(value, int)
        and not isinstance(value, bool)
        and len(str(abs(value))) > MAX_XLSX_SIGNIFICANT_DIGITS
    ):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def _resolve_aliases(links: Mapping[int, int | None]) -> dict[int, int]:
    resolved: dict[int, int] = {}
    for start in links:
        if start in resolved:
            continue
        path: list[int] = []
        positions: set[int] = set()
        current = start
        while current not in resolved:
            if current in positions:
                raise ValueError("inventory_alias_cycle")
            positions.add(current)
            path.append(current)
            target = links.get(current)
            if target is None:
                resolved[current] = current
                break
            current = target
        canonical_id = resolved[current]
        for alias_id in reversed(path):
            resolved[alias_id] = canonical_id
    return resolved


def _row_count(db: Session, statement) -> int:
    count_statement = select(func.count()).select_from(statement.order_by(None).subquery())
    return int(db.scalar(count_statement) or 0)


def _preflight_row_counts(db: Session, statements) -> int:
    total = 0
    for _title, _headers, statement in statements:
        total += _row_count(db, statement)
        if total > MAX_EXPORT_ROWS:
            raise OverflowError("inventory_export_too_large")
    return total


def _load_aliases(db: Session, *, remaining: int) -> dict[int, int]:
    links: dict[int, int | None] = {}
    statement = (
        select(InventoryCatalogPart.id, InventoryCatalogPart.merged_into_part_id)
        .order_by(InventoryCatalogPart.id)
        .limit(remaining + 1)
    )
    for part_id, target_id in db.execute(statement):
        links[part_id] = target_id
        if len(links) > remaining:
            raise OverflowError("inventory_export_too_large")
    return _resolve_aliases(links)


def _expanded_xlsx_rows(raw_row) -> Iterator[list]:
    values = list(raw_row)
    part_counts = [
        max(1, (len(value) + XLSX_TEXT_CHUNK_SIZE - 1) // XLSX_TEXT_CHUNK_SIZE)
        if isinstance(value, str)
        else 1
        for value in values
    ]
    part_count = max(part_counts, default=1)
    for part_index in range(part_count):
        row = []
        for value, value_part_count in zip(values, part_counts, strict=True):
            if isinstance(value, str):
                offset = part_index * XLSX_TEXT_CHUNK_SIZE
                chunk = value[offset : offset + XLSX_TEXT_CHUNK_SIZE]
                row.append(_safe_cell(chunk) if part_index < value_part_count else None)
            else:
                row.append(_xlsx_cell(value) if part_index == 0 else None)
        row.extend((part_index + 1, part_count))
        yield row


def _build_xlsx(db: Session, statements, *, canonical_ids: dict[int, int]):
    # Ownership is transferred to _file_chunks, which closes after streaming.
    spool = tempfile.SpooledTemporaryFile(  # noqa: SIM115
        max_size=XLSX_SPOOL_MEMORY_BYTES, mode="w+b"
    )
    book = Workbook(write_only=True)
    sheets = []
    for title, headers, _statement in statements:
        sheet = book.create_sheet(title)
        sheet.append((*headers, *CONTINUATION_HEADERS))
        sheets.append(sheet)

    overflow = False
    actual_rows = 0
    for (title, headers, statement), sheet in zip(statements, sheets, strict=True):
        sheet_rows = 0
        for raw_row in db.execute(statement.execution_options(yield_per=1_000)):
            row = list(raw_row)
            if title == "Движения" and row[3] is not None:
                row[4] = canonical_ids.get(row[3], row[3])
            elif title == "Поставки" and row[13] is not None:
                row[14] = canonical_ids.get(row[13], row[13])
            elif title == "Инвентаризации" and row[10] is not None:
                row[11] = canonical_ids.get(row[10], row[10])
            for expanded in _expanded_xlsx_rows(row):
                actual_rows += 1
                if len(canonical_ids) + actual_rows > MAX_EXPORT_ROWS:
                    overflow = True
                    break
                sheet.append(expanded)
                sheet_rows += 1
            if overflow:
                break
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers) + len(CONTINUATION_HEADERS))}{sheet_rows + 1}"
        )
        if overflow:
            break

    if overflow:
        # Saving finalizes write-only worksheet XML and prevents leaked generators.
        book.save(spool)
        spool.close()
        raise OverflowError("inventory_export_too_large")
    book.save(spool)
    spool.seek(0)
    return spool


def _file_chunks(file_obj) -> Iterator[bytes]:
    try:
        while chunk := file_obj.read(64 * 1024):
            yield chunk
    finally:
        file_obj.close()


def build_inventory_export(
    db: Session,
    user: User,
    *,
    park_id: int | None,
    all_parks: bool,
    format: str,
) -> tuple[str, str, Iterator[bytes]]:
    park_ids = _park_ids(db, user, park_id=park_id, all_parks=all_parks)
    if format not in {"csv", "xlsx"}:
        raise ValueError("inventory_export_format_invalid")
    stock_statement = _stock_statement(park_ids)
    if format == "csv":
        if _row_count(db, stock_statement) > MAX_EXPORT_ROWS:
            raise OverflowError("inventory_export_too_large")
    else:
        statements = [
            ("Остатки", STOCK_HEADERS, stock_statement),
            ("Движения", MOVEMENT_HEADERS, _movement_statement(park_ids)),
            ("Поставки", RECEIPT_HEADERS, _receipt_statement(park_ids)),
            ("Инвентаризации", COUNT_HEADERS, _count_statement(park_ids)),
        ]
        row_count = _preflight_row_counts(db, statements)
        canonical_ids = _load_aliases(db, remaining=MAX_EXPORT_ROWS - row_count)
        spool = _build_xlsx(
            db,
            statements,
            canonical_ids=canonical_ids,
        )
    suffix = "all" if all_parks else f"park-{park_ids[0]}"
    if format == "csv":
        return (
            "text/csv; charset=utf-8",
            f"Склад-{suffix}.csv",
            _csv_chunks(db.execute(stock_statement.execution_options(yield_per=1_000))),
        )
    return XLSX_MEDIA_TYPE, f"Склад-{suffix}.xlsx", _file_chunks(spool)
